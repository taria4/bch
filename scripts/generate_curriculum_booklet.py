#!/usr/bin/env python3
"""Generate the Persian curriculum outline as DOCX, HTML, PDF, and infographics."""

from __future__ import annotations

import html
import os
import re
import signal
import subprocess
import tempfile
import time
from pathlib import Path

import markdown
from PIL import Image, ImageDraw, ImageFont, features
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "computer-skills-curriculum-outline.md"
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "output"
COVER = ASSETS / "computer-skills-curriculum-cover.jpg"
FONT_REGULAR = Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf")
FONT_BOLD = Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Bold.ttf")

BLUE = "#075985"
GREEN = "#166534"
VIOLET = "#6B21A8"
ORANGE = "#C2410C"
DARK = "#1F2937"
LIGHT = "#F3F4F6"


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size=size)


def rtl_text(draw: ImageDraw.ImageDraw, xy, text: str, *, fnt, fill, anchor="mm"):
    kwargs = {"font": fnt, "fill": fill, "anchor": anchor}
    if features.check("raqm"):
        kwargs.update({"direction": "rtl", "language": "fa"})
    draw.text(xy, text, **kwargs)


def rounded_card(draw, box, fill, outline, width=4, radius=28):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def create_infographics() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    regular = font(FONT_REGULAR, 38)
    small = font(FONT_REGULAR, 30)
    bold = font(FONT_BOLD, 50)

    # Three-grade pathway
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    rtl_text(draw, (800, 80), "مسیر رشد سه‌پایه‌ای", fnt=font(FONT_BOLD, 62), fill=DARK)
    cards = [
        ((1050, 190, 1510, 730), "#E0F2FE", BLUE, "پایه هفتم", "می‌آموزم", ["آشنایی", "تمرین هدایت‌شده", "محصول کوتاه"]),
        ((570, 190, 1030, 730), "#DCFCE7", GREEN, "پایه هشتم", "حل می‌کنم", ["ترکیب مهارت‌ها", "حل مسئله", "محصول کاربردی"]),
        ((90, 190, 550, 730), "#F3E8FF", VIOLET, "پایه نهم", "می‌سازم", ["طراحی مستقل", "آزمون و اصلاح", "محصول کامل"]),
    ]
    for box, bg, color, grade, slogan, lines in cards:
        rounded_card(draw, box, bg, color, width=6)
        cx = (box[0] + box[2]) // 2
        rtl_text(draw, (cx, 265), grade, fnt=bold, fill=color)
        rtl_text(draw, (cx, 360), slogan, fnt=font(FONT_BOLD, 60), fill=DARK)
        for index, line in enumerate(lines):
            rtl_text(draw, (cx, 475 + index * 72), f"— {line}", fnt=regular, fill=DARK)
    rtl_text(draw, (800, 820), "از اجرای گام‌به‌گام تا تولید محصول مستقل", fnt=regular, fill=DARK)
    image.save(ASSETS / "infographic-three-grade-path.png", dpi=(180, 180))

    # Problem-solving cycle
    image = Image.new("RGB", (1600, 900), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    rtl_text(draw, (800, 80), "چرخه حل مسئله دیجیتال", fnt=font(FONT_BOLD, 62), fill=DARK)
    stages = [
        ("۱", "تعریف مسئله", BLUE),
        ("۲", "جمع‌آوری اطلاعات", GREEN),
        ("۳", "طراحی راه‌حل", VIOLET),
        ("۴", "اجرا", ORANGE),
        ("۵", "آزمون و اصلاح", "#B91C1C"),
        ("۶", "ارائه و بازتاب", "#0F766E"),
    ]
    positions = [(1250, 260), (800, 210), (350, 260), (350, 610), (800, 660), (1250, 610)]
    for idx, ((number, label, color), (cx, cy)) in enumerate(zip(stages, positions)):
        if idx < len(stages) - 1:
            nx, ny = positions[idx + 1]
        else:
            nx, ny = positions[0]
        draw.line((cx, cy, nx, ny), fill="#94A3B8", width=10)
    for (number, label, color), (cx, cy) in zip(stages, positions):
        rounded_card(draw, (cx - 185, cy - 82, cx + 185, cy + 82), "white", color, width=6)
        rtl_text(draw, (cx + 135, cy), number, fnt=font(FONT_BOLD, 44), fill=color)
        rtl_text(draw, (cx - 15, cy), label, fnt=small, fill=DARK)
    rtl_text(draw, (800, 820), "هر محصول خوب، دست‌کم یک بار آزمایش و اصلاح می‌شود.", fnt=regular, fill=DARK)
    image.save(ASSETS / "infographic-problem-solving-cycle.png", dpi=(180, 180))

    # Tool integration map
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    rtl_text(draw, (800, 70), "پیوند ابزارها در پروژه «مدرسه مجازی»", fnt=font(FONT_BOLD, 56), fill=DARK)
    center = (800, 450)
    draw.ellipse((590, 330, 1010, 570), fill="#E0F2FE", outline=BLUE, width=7)
    rtl_text(draw, center, "مدرسه مجازی", fnt=font(FONT_BOLD, 52), fill=BLUE)
    items = [
        ("Word", "کتابچه", (250, 220), BLUE),
        ("Excel", "کارنامه و فاکتور", (800, 180), GREEN),
        ("PowerPoint", "معرفی مدرسه", (1350, 220), ORANGE),
        ("Scratch", "بازی آموزشی", (250, 680), VIOLET),
        ("App Inventor", "اپلیکیشن", (800, 730), "#0F766E"),
        ("Python + AI", "محاسبه و رسانه", (1350, 680), "#B91C1C"),
    ]
    for title, subtitle, (cx, cy), color in items:
        draw.line((center[0], center[1], cx, cy), fill="#CBD5E1", width=8)
        rounded_card(draw, (cx - 195, cy - 90, cx + 195, cy + 90), LIGHT, color, width=5)
        draw.text((cx, cy - 24), title, font=font(FONT_BOLD, 35), fill=color, anchor="mm")
        rtl_text(draw, (cx, cy + 36), subtitle, fnt=small, fill=DARK)
    image.save(ASSETS / "infographic-tool-integration.png", dpi=(180, 180))


def body_markdown() -> str:
    text = SOURCE.read_text(encoding="utf-8")
    parts = text.split("\n---\n", 1)
    return parts[1] if len(parts) == 2 else text


def build_html() -> Path:
    md = body_markdown()
    content = markdown.markdown(md, extensions=["tables", "fenced_code", "sane_lists"])
    # Resolve all repository-relative image paths for Chrome.
    content = content.replace('src="../assets/', f'src="file://{ASSETS}/')
    cover_url = COVER.as_uri()
    css = f"""
    @font-face {{
      font-family: NotoArabic;
      src: url('file://{FONT_REGULAR}');
      font-weight: 400;
    }}
    @font-face {{
      font-family: NotoArabic;
      src: url('file://{FONT_BOLD}');
      font-weight: 700;
    }}
    @page {{ size: A4; margin: 18mm 17mm 19mm 20mm; }}
    * {{ box-sizing: border-box; }}
    body {{
      direction: rtl; font-family: NotoArabic, sans-serif; color: {DARK};
      font-size: 11pt; line-height: 1.75; margin: 0;
    }}
    .cover {{
      height: 250mm; page-break-after: always; break-inside: avoid; display: flex; flex-direction: column;
      justify-content: space-between; text-align: center; padding: 8mm 0;
    }}
    .cover h1 {{ color: {BLUE}; font-size: 29pt; line-height: 1.45; margin: 0; page-break-before: auto; border: 0; padding: 0; }}
    .cover h2 {{ font-size: 17pt; color: {DARK}; margin: 2mm 0; }}
    .cover img {{ width: 100%; border-radius: 5mm; margin: 8mm 0; }}
    .grade-strip {{ display: flex; gap: 4mm; direction: rtl; }}
    .grade {{ flex: 1; color: white; border-radius: 3mm; padding: 4mm 2mm; font-weight: 700; }}
    h1 {{ color: {BLUE}; font-size: 20pt; page-break-before: always; margin-top: 0; border-bottom: 2px solid {BLUE}; padding-bottom: 3mm; }}
    h2 {{ color: {VIOLET}; font-size: 15pt; break-after: avoid; margin-top: 8mm; }}
    h3 {{ color: {GREEN}; font-size: 12.5pt; break-after: avoid; }}
    p {{ margin: 2mm 0 3mm; }}
    ul, ol {{ margin: 2mm 0 4mm; padding-right: 7mm; }}
    li {{ margin-bottom: 1mm; }}
    blockquote {{
      border-right: 4px solid {ORANGE}; background: #FFF7ED; margin: 5mm 0;
      padding: 4mm 5mm; border-radius: 2mm; break-inside: avoid;
    }}
    table {{ width: 100%; border-collapse: collapse; margin: 5mm 0; font-size: 9.2pt; }}
    thead {{ display: table-header-group; }}
    tr {{ break-inside: avoid; }}
    th {{ background: {BLUE}; color: white; font-weight: 700; }}
    th, td {{ border: 1px solid #CBD5E1; padding: 2.2mm; text-align: right; vertical-align: top; }}
    tr:nth-child(even) td {{ background: #F8FAFC; }}
    img {{ max-width: 100%; height: auto; display: block; margin: 6mm auto 2mm; border-radius: 2mm; break-inside: avoid; }}
    p:has(> img) {{ color: #475569; font-size: 9pt; text-align: center; break-inside: avoid; }}
    code {{ direction: ltr; unicode-bidi: embed; font-family: monospace; background: #F1F5F9; padding: 0 1mm; }}
    strong {{ font-weight: 700; }}
    hr {{ border: 0; border-top: 1px solid #CBD5E1; margin: 7mm 0; }}
    .footer-note {{ color: #64748B; font-size: 9pt; text-align: center; }}
    """
    page = f"""<!doctype html>
<html lang="fa" dir="rtl"><head><meta charset="utf-8"><title>طرح تفصیلی مجموعه آموزش مهارت‌های کامپیوتری</title><style>{css}</style></head>
<body>
<section class="cover">
  <div>
    <h1>طرح تفصیلی مجموعه آموزش مهارت‌های کامپیوتری</h1>
    <h2>ویژه دانش‌آموزان پایه‌های هفتم، هشتم و نهم</h2>
  </div>
  <img src="{cover_url}" alt="دانش‌آموزان در حال یادگیری مهارت‌های دیجیتال">
  <div class="grade-strip">
    <div class="grade" style="background:{BLUE}">هفتم<br>می‌آموزم</div>
    <div class="grade" style="background:{GREEN}">هشتم<br>حل می‌کنم</div>
    <div class="grade" style="background:{VIOLET}">نهم<br>می‌سازم</div>
  </div>
  <p class="footer-note">نسخه ۱.۰ — سال تحصیلی ۱۴۰۵–۱۴۰۶</p>
</section>
{content}
</body></html>"""
    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / "computer-skills-curriculum-outline.html"
    target.write_text(page, encoding="utf-8")
    return target


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_rtl(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def style_run(run, *, bold: bool | None = None, size: float = 11, color: str = DARK) -> None:
    run.font.name = "Noto Sans Arabic"
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.removeprefix("#"))
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), "Noto Sans Arabic")
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip().replace("  ", " ")


def add_text_with_markup(paragraph, text: str, size: float = 11, color: str = DARK) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if is_bold else part[1:-1] if is_code else part
        run = paragraph.add_run(value)
        style_run(run, bold=is_bold, size=size, color=color)
        if is_code:
            run.font.name = "DejaVu Sans Mono"


def add_page_number(paragraph) -> None:
    set_rtl(paragraph)
    run = paragraph.add_run("صفحه ")
    style_run(run, size=8, color="#64748B")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_docx_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("طرح تفصیلی مجموعه\nآموزش مهارت‌های کامپیوتری")
    style_run(run, bold=True, size=26, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ویژه دانش‌آموزان پایه‌های هفتم، هشتم و نهم")
    style_run(run, bold=True, size=16)
    doc.add_picture(str(COVER), width=Inches(6.65))
    table = doc.add_table(rows=1, cols=3)
    labels = [("هفتم\nمی‌آموزم", BLUE), ("هشتم\nحل می‌کنم", GREEN), ("نهم\nمی‌سازم", VIOLET)]
    for cell, (label, color) in zip(table.rows[0].cells, labels):
        set_cell_shading(cell, color.removeprefix("#"))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        style_run(run, bold=True, size=12, color="#FFFFFF")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\nنسخه ۱.۰ — سال تحصیلی ۱۴۰۵–۱۴۰۶")
    style_run(run, size=10, color="#64748B")
    doc.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [clean_inline(cell) for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row):
            rows.append(row)
        index += 1
    return rows, index


def build_docx() -> Path:
    print("Building DOCX...", flush=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.8)
    add_docx_cover(doc)

    header = section.header.paragraphs[0]
    set_rtl(header)
    run = header.add_run("مجموعه آموزش مهارت‌های کامپیوتری | طرح تفصیلی")
    style_run(run, bold=True, size=8.5, color=BLUE)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    lines = body_markdown().splitlines()
    index = 0
    while index < len(lines):
        if index % 50 == 0:
            print(f"DOCX source line {index}/{len(lines)}", flush=True)
        raw = lines[index]
        line = raw.strip()
        if not line or line == "---":
            index += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            alt, rel = image_match.groups()
            image_path = (SOURCE.parent / rel).resolve()
            if image_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(image_path), width=Inches(6.45))
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(alt)
                style_run(run, size=8.5, color="#64748B")
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            if level == 1 and index > 0:
                doc.add_page_break()
            p = doc.add_paragraph()
            set_rtl(p)
            p.paragraph_format.keep_with_next = True
            sizes = {1: 20, 2: 15, 3: 12.5}
            colors = {1: BLUE, 2: VIOLET, 3: GREEN}
            run = p.add_run(text)
            style_run(run, bold=True, size=sizes[level], color=colors[level])
            index += 1
            continue

        if line.startswith("|"):
            rows, next_index = parse_table(lines, index)
            if rows:
                cols = max(len(row) for row in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                table.autofit = True
                table_pr = table._tbl.tblPr
                bidi_visual = OxmlElement("w:bidiVisual")
                table_pr.append(bidi_visual)
                for row_index, row in enumerate(rows):
                    for col_index in range(cols):
                        cell = table.cell(row_index, col_index)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        p = cell.paragraphs[0]
                        set_rtl(p)
                        value = row[col_index] if col_index < len(row) else ""
                        run = p.add_run(value)
                        style_run(run, bold=row_index == 0, size=8.6, color="#FFFFFF" if row_index == 0 else DARK)
                        if row_index == 0:
                            set_cell_shading(cell, BLUE.removeprefix("#"))
                        elif row_index % 2:
                            set_cell_shading(cell, "F8FAFC")
            index = next_index
            continue

        if line.startswith(">"):
            p = doc.add_paragraph()
            set_rtl(p)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.right_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "FFF7ED")
            p_pr.append(shd)
            add_text_with_markup(p, line[1:].strip(), size=10.5)
            index += 1
            continue

        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            set_rtl(p)
            p.paragraph_format.space_after = Pt(2)
            add_text_with_markup(p, re.sub(r"^[-*]\s+", "", line))
            index += 1
            continue

        # Keep Persian-numbered lines as ordinary RTL paragraphs.
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(5)
        add_text_with_markup(p, line)
        index += 1

    core = doc.core_properties
    core.title = "طرح تفصیلی مجموعه آموزش مهارت‌های کامپیوتری"
    core.subject = "پایه‌های هفتم، هشتم و نهم"
    core.comments = "نسخه ۱.۰"
    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / "computer-skills-curriculum-outline.docx"
    doc.save(target)
    return target


def build_pdf(html_path: Path) -> Path:
    target = OUTPUT / "computer-skills-curriculum-outline.pdf"
    target.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix="curriculum-pdf-") as profile:
        command = [
            "google-chrome",
            "--headless=new",
            "--no-sandbox",
            "--disable-gpu",
            f"--user-data-dir={profile}",
            "--allow-file-access-from-files",
            "--no-pdf-header-footer",
            f"--print-to-pdf={target}",
            html_path.as_uri(),
        ]
        process = subprocess.Popen(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            start_new_session=True,
        )
        last_size = -1
        stable_checks = 0
        for _ in range(300):
            if target.exists():
                size = target.stat().st_size
                stable_checks = stable_checks + 1 if size == last_size else 0
                last_size = size
                if stable_checks >= 10 and size > 10_000 and target.read_bytes()[:4] == b"%PDF":
                    break
            if process.poll() is not None:
                break
            time.sleep(0.1)
        if process.poll() is None:
            os.killpg(process.pid, signal.SIGTERM)
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid, signal.SIGKILL)
        if not target.exists() or target.stat().st_size < 10_000 or target.read_bytes()[:4] != b"%PDF":
            stderr = process.stderr.read() if process.stderr else ""
            raise RuntimeError(f"Chrome did not produce a valid PDF: {stderr}")
    return target


def main() -> None:
    print("Creating infographics...", flush=True)
    create_infographics()
    print("Building HTML...", flush=True)
    html_path = build_html()
    docx_path = build_docx()
    print("Building PDF...", flush=True)
    pdf_path = build_pdf(html_path)
    print(f"Generated: {docx_path.relative_to(ROOT)}")
    print(f"Generated: {pdf_path.relative_to(ROOT)}")
    print(f"Generated: {html_path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
