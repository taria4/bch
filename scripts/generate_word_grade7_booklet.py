#!/usr/bin/env python3
"""Generate the branded Grade 7 Microsoft Word pilot booklet and companion files."""

from __future__ import annotations

import os
import re
import signal
import subprocess
import tempfile
import time
from pathlib import Path

import markdown
from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

import generate_curriculum_booklet as base


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "word-grade7-booklet.md"
ASSETS = ROOT / "assets"
VISUALS = ASSETS / "word-grade7"
LOGO = ASSETS / "baqerololoum-logo.png"
OUTPUT = ROOT / "output" / "word-grade7"
FILES = OUTPUT / "files"

BLUE = "#075985"
SKY = "#E0F2FE"
TEAL = "#0F766E"
GREEN = "#166534"
VIOLET = "#6B21A8"
ORANGE = "#C2410C"
RED = "#B91C1C"
DARK = "#1F2937"
GRAY = "#64748B"
LIGHT = "#F8FAFC"


def body_markdown() -> str:
    text = SOURCE.read_text(encoding="utf-8")
    parts = text.split("\n---\n", 1)
    return parts[1] if len(parts) == 2 else text


def fnt(size: int, bold: bool = False):
    return base.font(base.FONT_BOLD if bold else base.FONT_REGULAR, size)


def text(draw, xy, value: str, size: int, color=DARK, bold=False, anchor="mm"):
    base.rtl_text(draw, xy, value, fnt=fnt(size, bold), fill=color, anchor=anchor)


def card(draw, box, fill="white", outline="#CBD5E1", width=4, radius=24):
    base.rounded_card(draw, box, fill, outline, width=width, radius=radius)


def number_badge(draw, x: int, y: int, number: int, color=BLUE):
    draw.ellipse((x - 30, y - 30, x + 30, y + 30), fill=color)
    draw.text((x, y), str(number), font=fnt(28, True), fill="white", anchor="mm")


def save_visual(image: Image.Image, name: str) -> None:
    image.save(VISUALS / name, dpi=(180, 180))


def visual_title(draw, title: str, subtitle: str | None = None) -> None:
    text(draw, (800, 65), title, 54, BLUE, True)
    if subtitle:
        text(draw, (800, 120), subtitle, 27, GRAY)


def create_visuals() -> None:
    VISUALS.mkdir(parents=True, exist_ok=True)

    # Cover illustration
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    card(draw, (250, 110, 1350, 790), SKY, BLUE, 8, 34)
    draw.rectangle((390, 215, 1210, 680), fill="white", outline=BLUE, width=7)
    draw.rectangle((390, 215, 1210, 285), fill=BLUE)
    for x, label in ((460, "File"), (600, "Home"), (750, "Insert"), (915, "Layout")):
        draw.text((x, 250), label, font=fnt(29, True), fill="white", anchor="mm")
    draw.rectangle((500, 350, 1100, 610), fill="#F8FAFC", outline="#CBD5E1", width=4)
    text(draw, (800, 410), "اولین سند حرفه‌ای من", 48, VIOLET, True)
    draw.line((580, 480, 1020, 480), fill="#94A3B8", width=8)
    draw.line((620, 535, 980, 535), fill="#CBD5E1", width=8)
    draw.line((680, 590, 920, 590), fill="#CBD5E1", width=8)
    text(draw, (800, 735), "تایپ • تصویر • جدول • صفحه‌آرایی", 34, DARK)
    save_visual(image, "00-cover-illustration.png")

    # 1. Interface
    image = Image.new("RGB", (1600, 900), LIGHT)
    draw = ImageDraw.Draw(image)
    visual_title(draw, "بخش‌های اصلی محیط Word", "طرح شماتیک؛ جای بعضی دکمه‌ها ممکن است متفاوت باشد")
    draw.rectangle((110, 170, 1490, 790), fill="white", outline=DARK, width=5)
    draw.rectangle((110, 170, 1490, 235), fill=BLUE)
    draw.rectangle((110, 235, 1490, 310), fill="#E2E8F0")
    draw.rectangle((110, 310, 1490, 410), fill="#F1F5F9")
    draw.rectangle((400, 450, 1200, 735), fill="white", outline="#94A3B8", width=3)
    draw.rectangle((110, 745, 1490, 790), fill="#E2E8F0")
    labels = [
        (1, 1270, 202, "نوار عنوان"),
        (2, 1300, 270, "زبانه‌ها"),
        (3, 1290, 360, "Ribbon"),
        (4, 240, 430, "خط‌کش"),
        (5, 800, 590, "صفحه سند"),
        (6, 250, 767, "نوار وضعیت"),
        (7, 1340, 767, "Zoom"),
    ]
    for number, x, y, label in labels:
        number_badge(draw, x, y, number)
        text(draw, (x - 75 if x > 800 else x + 90, y), label, 26)
    save_visual(image, "01-word-interface.png")

    # 2. Save workflow
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    visual_title(draw, "مسیر ذخیره نخستین سند")
    steps = ["سند تازه", "Save As", "انتخاب پوشه", "نام فایل", "Save"]
    colors = [BLUE, VIOLET, TEAL, ORANGE, GREEN]
    centers = [1350, 1075, 800, 525, 250]
    for index, (label, color, cx) in enumerate(zip(steps, colors, centers), 1):
        card(draw, (cx - 105, 300, cx + 105, 590), "#F8FAFC", color, 5)
        number_badge(draw, cx, 355, index, color)
        text(draw, (cx, 470), label, 34, color, True)
        if index < len(steps):
            draw.line((cx - 110, 445, cx - 160, 445), fill="#94A3B8", width=9)
            draw.polygon([(cx - 180, 445), (cx - 150, 425), (cx - 150, 465)], fill="#94A3B8")
    text(draw, (800, 700), "میان‌بر ذخیره: Ctrl + S", 38, DARK, True)
    save_visual(image, "02-save-workflow.png")

    # 3. Persian typing
    image = Image.new("RGB", (1600, 900), LIGHT)
    draw = ImageDraw.Draw(image)
    visual_title(draw, "آماده‌سازی Word برای تایپ فارسی")
    card(draw, (160, 190, 1440, 710), "white", "#CBD5E1", 4)
    labels = [
        ("۱", "زبان فارسی", "Windows + Space"),
        ("۲", "جهت متن", "Right-to-Left"),
        ("۳", "مکان‌نما", "محل شروع تایپ"),
        ("۴", "نیم‌فاصله", "Ctrl + Shift + 2"),
        ("۵", "سطر تازه", "Enter"),
        ("۶", "حذف", "Backspace"),
    ]
    positions = [(1180, 300), (800, 300), (420, 300), (1180, 535), (800, 535), (420, 535)]
    for (num, title, subtitle), (cx, cy) in zip(labels, positions):
        card(draw, (cx - 160, cy - 85, cx + 160, cy + 85), SKY, BLUE, 3)
        draw.text((cx + 120, cy - 35), num, font=fnt(34, True), fill=BLUE, anchor="mm")
        text(draw, (cx, cy - 25), title, 31, DARK, True)
        draw.text((cx, cy + 38), subtitle, font=fnt(24), fill=GRAY, anchor="mm")
    text(draw, (800, 790), "دانش‌آموزان می‌توانند درست و خوانا تایپ کنند.", 36, VIOLET, True)
    save_visual(image, "03-persian-typing.png")

    # 4. Editing tools
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    visual_title(draw, "انتخاب، ویرایش و جست‌وجوی متن")
    tools = [("Copy", "Ctrl+C"), ("Cut", "Ctrl+X"), ("Paste", "Ctrl+V"), ("Undo", "Ctrl+Z"), ("Find", "Ctrl+F"), ("Replace", "Ctrl+H")]
    positions = [(1250, 300), (800, 300), (350, 300), (1250, 590), (800, 590), (350, 590)]
    for index, ((name, shortcut), (cx, cy)) in enumerate(zip(tools, positions), 1):
        card(draw, (cx - 180, cy - 100, cx + 180, cy + 100), LIGHT, [BLUE, VIOLET, TEAL, ORANGE, GREEN, RED][index - 1], 5)
        number_badge(draw, cx + 140, cy - 60, index, [BLUE, VIOLET, TEAL, ORANGE, GREEN, RED][index - 1])
        draw.text((cx, cy - 20), name, font=fnt(42, True), fill=DARK, anchor="mm")
        draw.text((cx, cy + 48), shortcut, font=fnt(30), fill=GRAY, anchor="mm")
    save_visual(image, "04-editing-tools.png")

    # 5. Formatting
    image = Image.new("RGB", (1600, 900), LIGHT)
    draw = ImageDraw.Draw(image)
    visual_title(draw, "متن خام و متن قالب‌بندی‌شده")
    card(draw, (150, 190, 750, 730), "white", "#94A3B8", 4)
    card(draw, (850, 190, 1450, 730), "white", BLUE, 5)
    text(draw, (450, 255), "قبل", 40, GRAY, True)
    text(draw, (1150, 255), "بعد", 40, BLUE, True)
    text(draw, (450, 370), "عنوان درس", 34)
    text(draw, (450, 455), "این یک متن ساده و بدون نظم است.", 27)
    text(draw, (1150, 350), "عنوان درس", 48, VIOLET, True)
    text(draw, (1150, 445), "این متن خوانا و هماهنگ است.", 31)
    features = [("۱", "Font"), ("۲", "Size"), ("۳", "Bold"), ("۴", "Color"), ("۵", "Highlight"), ("۶", "Alignment")]
    for index, (num, label) in enumerate(features):
        x = 1320 - (index % 3) * 170
        y = 570 + (index // 3) * 85
        number_badge(draw, x, y, int(num), BLUE)
        draw.text((x - 75, y), label, font=fnt(22), fill=DARK, anchor="mm")
    save_visual(image, "05-text-formatting.png")

    # 6. Paragraph
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    visual_title(draw, "تنظیم‌های مهم پاراگراف")
    draw.rectangle((280, 180, 1320, 760), fill=LIGHT, outline="#94A3B8", width=4)
    for y, width in ((290, 760), (360, 820), (445, 700), (535, 780)):
        draw.line((1260 - width, y, 1260, y), fill=DARK, width=9)
    draw.line((1280, 270, 1280, 560), fill=BLUE, width=5)
    draw.line((1240, 270, 1320, 270), fill=BLUE, width=5)
    draw.line((1240, 560, 1320, 560), fill=BLUE, width=5)
    callouts = [(1, 1220, 650, "راست‌چین"), (2, 940, 650, "فاصله خطوط"), (3, 650, 650, "فاصله بند"), (4, 370, 650, "تورفتگی")]
    for number, x, y, label in callouts:
        number_badge(draw, x, y, number)
        text(draw, (x, y + 55), label, 26)
    save_visual(image, "06-paragraph-layout.png")

    # 7. Picture layout
    image = Image.new("RGB", (1600, 900), LIGHT)
    draw = ImageDraw.Draw(image)
    visual_title(draw, "درج و تنظیم درست تصویر")
    draw.rectangle((300, 180, 1300, 760), fill="white", outline="#94A3B8", width=4)
    draw.rectangle((560, 320, 1040, 620), fill=SKY, outline=BLUE, width=6)
    text(draw, (800, 470), "تصویر آموزشی", 44, BLUE, True)
    for x, y in ((560, 320), (1040, 320), (560, 620), (1040, 620)):
        draw.rectangle((x - 12, y - 12, x + 12, y + 12), fill="white", outline=BLUE, width=4)
    labels = [(1, 1210, 245, "Insert Picture"), (2, 1110, 660, "دستگیره گوشه"), (3, 800, 700, "حفظ تناسب"), (4, 430, 470, "Wrap Text"), (5, 800, 790, "شرح تصویر")]
    for number, x, y, label in labels:
        number_badge(draw, x, y, number, VIOLET)
        text(draw, (x, y + 55 if y < 750 else y), label, 24, DARK)
    save_visual(image, "07-picture-layout.png")

    # 8. Table
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    visual_title(draw, "اجزای یک جدول ساده")
    left, top, right, bottom = 260, 200, 1340, 720
    rows, cols = 5, 4
    cell_w, cell_h = (right - left) // cols, (bottom - top) // rows
    draw.rectangle((left, top, right, bottom), fill="white", outline=BLUE, width=6)
    draw.rectangle((left, top, right, top + cell_h), fill=SKY)
    for r in range(1, rows):
        draw.line((left, top + r * cell_h, right, top + r * cell_h), fill="#64748B", width=3)
    for c in range(1, cols):
        draw.line((left + c * cell_w, top, left + c * cell_w, bottom), fill="#64748B", width=3)
    for number, x, y, label in [
        (1, 1420, 250, "جدول"),
        (2, 1420, 430, "سطر"),
        (3, 1120, 785, "ستون"),
        (4, 800, 460, "سلول"),
        (5, 510, 250, "سرستون"),
        (6, 220, 660, "افزودن سطر"),
    ]:
        number_badge(draw, x, y, number, TEAL)
        text(draw, (x, y + 55), label, 24)
    save_visual(image, "08-table-parts.png")

    # 9. Page layout
    image = Image.new("RGB", (1600, 900), LIGHT)
    draw = ImageDraw.Draw(image)
    visual_title(draw, "صفحه‌آرایی استاندارد")
    draw.rectangle((480, 160, 1120, 790), fill="white", outline=DARK, width=5)
    draw.rectangle((550, 230, 1050, 720), outline="#94A3B8", width=3)
    draw.rectangle((550, 230, 1050, 290), fill=SKY)
    draw.rectangle((550, 660, 1050, 720), fill="#F1F5F9")
    labels = [(1, 1260, 200, "A4"), (2, 1260, 380, "Portrait"), (3, 1260, 590, "Margins"), (4, 380, 250, "Header"), (5, 380, 675, "Footer"), (6, 800, 755, "Page Number")]
    for number, x, y, label in labels:
        number_badge(draw, x, y, number, ORANGE)
        text(draw, (x, y + 55), label, 24)
    save_visual(image, "09-page-layout.png")

    # 10. Print and PDF
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    visual_title(draw, "کنترل پیش از چاپ و خروجی PDF")
    steps = [
        ("۱", "Save", "ذخیره آخرین تغییر"),
        ("۲", "Print Preview", "کنترل همه صفحه‌ها"),
        ("۳", "Settings", "چاپگر و محدوده"),
        ("۴", "Export PDF", "ساخت فایل نهایی"),
        ("۵", "Review", "بازکردن و بازبینی"),
    ]
    for index, (num, title, subtitle) in enumerate(steps):
        cx = 1360 - index * 280
        card(draw, (cx - 120, 280, cx + 120, 620), LIGHT, [BLUE, VIOLET, TEAL, ORANGE, GREEN][index], 5)
        number_badge(draw, cx, 345, int(num), [BLUE, VIOLET, TEAL, ORANGE, GREEN][index])
        draw.text((cx, 455), title, font=fnt(31, True), fill=DARK, anchor="mm")
        text(draw, (cx, 540), subtitle, 23, GRAY)
        if index < len(steps) - 1:
            draw.line((cx - 125, 450, cx - 155, 450), fill="#94A3B8", width=8)
    text(draw, (800, 740), "DOCX برای ویرایش  |  PDF برای مشاهده و چاپ", 36, BLUE, True)
    save_visual(image, "10-print-pdf.png")

    # 11. Final project
    image = Image.new("RGB", (1600, 900), LIGHT)
    draw = ImageDraw.Draw(image)
    visual_title(draw, "نقشه اجرای پروژه پایانی Word")
    steps = ["تحلیل", "پوشه‌بندی", "تایپ", "قالب‌بندی", "تصویر و جدول", "صفحه‌آرایی", "بازبینی", "تحویل"]
    colors = [BLUE, VIOLET, TEAL, ORANGE, GREEN, RED, "#0369A1", "#4338CA"]
    for index, (label, color) in enumerate(zip(steps, colors)):
        row, col = divmod(index, 4)
        cx = 1320 - col * 350
        cy = 305 + row * 300
        card(draw, (cx - 140, cy - 85, cx + 140, cy + 85), "white", color, 5)
        number_badge(draw, cx + 100, cy - 50, index + 1, color)
        text(draw, (cx, cy + 15), label, 31, DARK, True)
    text(draw, (800, 820), "محصول نهایی: دو صفحه کتاب درسی در قالب DOCX و PDF", 35, BLUE, True)
    save_visual(image, "11-final-project-map.png")

    # 12. Water-cycle illustration for sample project
    image = Image.new("RGB", (1600, 900), SKY)
    draw = ImageDraw.Draw(image)
    visual_title(draw, "چرخه آب در طبیعت")
    draw.ellipse((150, 170, 350, 370), fill="#FBBF24", outline=ORANGE, width=5)
    draw.polygon([(450, 680), (800, 260), (1150, 680)], fill="#94A3B8")
    draw.polygon([(690, 390), (800, 260), (910, 390)], fill="white")
    draw.ellipse((1060, 180, 1400, 330), fill="white", outline="#CBD5E1", width=4)
    draw.rectangle((0, 680, 1600, 900), fill="#38BDF8")
    arrows = [
        ((300, 650), (420, 400), "تبخیر"),
        ((1150, 360), (1250, 500), "بارش"),
        ((1050, 720), (650, 720), "جریان آب"),
    ]
    for (x1, y1), (x2, y2), label in arrows:
        draw.line((x1, y1, x2, y2), fill=BLUE, width=12)
        text(draw, ((x1 + x2) // 2, (y1 + y2) // 2 - 35), label, 32, BLUE, True)
    save_visual(image, "12-water-cycle.png")


def print_pdf(html_path: Path, target: Path) -> None:
    target.unlink(missing_ok=True)
    with tempfile.TemporaryDirectory(prefix="word-grade7-pdf-") as profile:
        command = [
            "google-chrome",
            "--headless=new",
            "--no-sandbox",
            "--disable-gpu",
            f"--user-data-dir={profile}",
            "--allow-file-access-from-files",
            "--no-pdf-header-footer",
            f"--print-to-pdf={target}",
            html_path.as_uri(),
        ]
        process = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, start_new_session=True)
        last_size, stable = -1, 0
        for _ in range(400):
            if target.exists():
                size = target.stat().st_size
                stable = stable + 1 if size == last_size else 0
                last_size = size
                if stable >= 10 and size > 10_000 and target.read_bytes()[:4] == b"%PDF":
                    break
            if process.poll() is not None:
                break
            time.sleep(0.1)
        if process.poll() is None:
            os.killpg(process.pid, signal.SIGTERM)
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid, signal.SIGKILL)
        if not target.exists() or target.stat().st_size < 10_000 or target.read_bytes()[:4] != b"%PDF":
            error = process.stderr.read() if process.stderr else ""
            raise RuntimeError(f"PDF generation failed: {error}")


def html_css() -> str:
    return f"""
    @font-face {{ font-family:NotoArabic; src:url('file://{base.FONT_REGULAR}'); font-weight:400; }}
    @font-face {{ font-family:NotoArabic; src:url('file://{base.FONT_BOLD}'); font-weight:700; }}
    @page {{ size:A4; margin:17mm 16mm 18mm 19mm; }}
    * {{ box-sizing:border-box; }}
    body {{ direction:rtl; font-family:NotoArabic,sans-serif; color:{DARK}; font-size:11pt; line-height:1.75; margin:0; }}
    .cover {{ height:252mm; page-break-after:always; break-inside:avoid; display:flex; flex-direction:column; justify-content:space-between; text-align:center; padding:2mm 0; }}
    .cover-logo {{ width:38mm; max-height:38mm; object-fit:contain; margin:0 auto; }}
    .cover-hero {{ width:100%; max-height:115mm; object-fit:contain; border-radius:4mm; }}
    .cover h1 {{ color:{BLUE}; font-size:29pt; line-height:1.4; margin:0; page-break-before:auto; border:0; padding:0; }}
    .cover h2 {{ color:{VIOLET}; font-size:18pt; margin:1mm 0; }}
    .meta {{ border-top:2px solid {BLUE}; padding-top:4mm; font-size:11pt; }}
    h1 {{ color:{BLUE}; font-size:20pt; page-break-before:always; border-bottom:2px solid {BLUE}; padding-bottom:3mm; margin-top:0; }}
    h2 {{ color:{VIOLET}; font-size:14.5pt; break-after:avoid; margin-top:7mm; }}
    h3 {{ color:{GREEN}; font-size:12.5pt; break-after:avoid; }}
    p {{ margin:2mm 0 3mm; }}
    ul,ol {{ margin:2mm 0 4mm; padding-right:7mm; }}
    li {{ margin-bottom:1mm; }}
    blockquote {{ border-right:4px solid {ORANGE}; background:#FFF7ED; margin:5mm 0; padding:4mm 5mm; border-radius:2mm; break-inside:avoid; }}
    table {{ width:100%; border-collapse:collapse; margin:5mm 0; font-size:9.3pt; }}
    thead {{ display:table-header-group; }}
    tr {{ break-inside:avoid; }}
    th {{ background:{BLUE}; color:white; }}
    th,td {{ border:1px solid #CBD5E1; padding:2.2mm; text-align:right; vertical-align:top; }}
    tr:nth-child(even) td {{ background:#F8FAFC; }}
    img {{ max-width:100%; height:auto; display:block; margin:5mm auto 1mm; border-radius:2mm; break-inside:avoid; }}
    p:has(> img) {{ text-align:center; font-size:9pt; color:{GRAY}; break-inside:avoid; }}
    code {{ direction:ltr; unicode-bidi:embed; font-family:monospace; background:#F1F5F9; padding:0 1mm; }}
    strong {{ font-weight:700; }}
    hr {{ border:0; border-top:1px solid #CBD5E1; margin:7mm 0; }}
    """


def build_html() -> Path:
    content = markdown.markdown(body_markdown(), extensions=["tables", "fenced_code", "sane_lists"])
    content = content.replace('src="../assets/', f'src="file://{ASSETS}/')
    page = f"""<!doctype html><html lang="fa" dir="rtl"><head><meta charset="utf-8">
<title>Microsoft Word پایه هفتم</title><style>{html_css()}</style></head><body>
<section class="cover">
  <img class="cover-logo" src="{LOGO.as_uri()}" alt="لوگوی مجتمع آموزشی فرهنگی باقرالعلوم (ع)">
  <div><h1>Microsoft Word ـ پایه هفتم</h1><h2>آموزش گام‌به‌گام واژه‌پردازی</h2></div>
  <img class="cover-hero" src="{(VISUALS / '00-cover-illustration.png').as_uri()}" alt="نماد آموزش Word">
  <div class="meta"><strong>مجتمع آموزشی فرهنگی باقرالعلوم (ع)</strong><br>
  نویسنده: مهندس نقی زاده<br>نسخه آزمایشی ۱.۰ — سال تحصیلی ۱۴۰۵–۱۴۰۶</div>
  <div>نام دانش‌آموز: .............................. &nbsp;&nbsp; کلاس: ...............</div>
</section>{content}</body></html>"""
    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / "word-grade7-booklet.html"
    target.write_text(page, encoding="utf-8")
    return target


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(1.45))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Microsoft Word ـ پایه هفتم")
    base.style_run(run, bold=True, size=26, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("آموزش گام‌به‌گام واژه‌پردازی")
    base.style_run(run, bold=True, size=17, color=VIOLET)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(VISUALS / "00-cover-illustration.png"), width=Inches(6.4))
    for value, bold in [
        ("مجتمع آموزشی فرهنگی باقرالعلوم (ع)", True),
        ("نویسنده: مهندس نقی زاده", True),
        ("نسخه آزمایشی ۱.۰ — سال تحصیلی ۱۴۰۵–۱۴۰۶", False),
        ("نام دانش‌آموز: ..............................     کلاس: ...............", False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        base.style_run(run, bold=bold, size=11, color=DARK if bold else GRAY)
    doc.add_page_break()


def build_docx() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.8)
    section.left_margin, section.right_margin = Cm(1.9), Cm(1.6)
    add_cover(doc)

    header = section.header.paragraphs[0]
    base.set_rtl(header)
    run = header.add_run("مجتمع آموزشی فرهنگی باقرالعلوم (ع) | Microsoft Word پایه هفتم")
    base.style_run(run, bold=True, size=8.2, color=BLUE)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.add_page_number(footer)

    lines = body_markdown().splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line == "---":
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            alt, rel = image_match.groups()
            path = (SOURCE.parent / rel).resolve()
            if path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Inches(6.45))
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(alt)
                base.style_run(run, size=8.5, color=GRAY)
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            if level == 1 and index > 0:
                doc.add_page_break()
            p = doc.add_paragraph()
            base.set_rtl(p)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(base.clean_inline(heading.group(2)))
            base.style_run(run, bold=True, size={1: 20, 2: 14.5, 3: 12.5}[level], color={1: BLUE, 2: VIOLET, 3: GREEN}[level])
            index += 1
            continue
        if line.startswith("|"):
            rows, next_index = base.parse_table(lines, index)
            if rows:
                cols = max(len(row) for row in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                bidi = OxmlElement("w:bidiVisual")
                table._tbl.tblPr.append(bidi)
                for r, row in enumerate(rows):
                    for c in range(cols):
                        cell = table.cell(r, c)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        p = cell.paragraphs[0]
                        base.set_rtl(p)
                        run = p.add_run(row[c] if c < len(row) else "")
                        base.style_run(run, bold=r == 0, size=8.5, color="#FFFFFF" if r == 0 else DARK)
                        if r == 0:
                            base.set_cell_shading(cell, BLUE.removeprefix("#"))
                        elif r % 2:
                            base.set_cell_shading(cell, "F8FAFC")
            index = next_index
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            base.set_rtl(p)
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = Cm(0.5)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "FFF7ED")
            p._p.get_or_add_pPr().append(shd)
            base.add_text_with_markup(p, line[1:].strip(), size=10.5)
            index += 1
            continue
        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            base.set_rtl(p)
            base.add_text_with_markup(p, re.sub(r"^[-*]\s+", "", line))
            index += 1
            continue
        p = doc.add_paragraph()
        base.set_rtl(p)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(5)
        base.add_text_with_markup(p, line)
        index += 1

    props = doc.core_properties
    props.title = "Microsoft Word پایه هفتم"
    props.author = "مهندس نقی زاده"
    props.subject = "مجتمع آموزشی فرهنگی باقرالعلوم (ع)"
    target = OUTPUT / "word-grade7-booklet.docx"
    doc.save(target)
    return target


def build_practice_file() -> Path:
    FILES.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    p = doc.add_paragraph()
    base.set_rtl(p)
    run = p.add_run("فایل تمرین خام ـ آزمون عملکردی Word پایه هفتم")
    base.style_run(run, bold=True, size=18, color=BLUE)
    p = doc.add_paragraph()
    base.set_rtl(p)
    base.add_text_with_markup(p, "این فایل عمداً بدون قالب‌بندی نهایی تهیه شده است. آن را مطابق دستور آزمون اصلاح کنید.")
    for value in [
        "انرژی خورشیدی",
        "خورشید یکی از مهم ترین منابع انرژی است .انسان می تواند با استفاده از صفحه های خورشیدی نور را به برق تبدیل کند.",
        "مزایای انرژی خورشیدی",
        "کاهش آلودگی هوا",
        "کمک به صرفه جویی در سوخت",
        "امکان استفاده در مناطق دور",
        "مراحل کار",
        "دریافت نور خورشید - تبدیل نور به برق - ذخیره یا مصرف انرژی",
        "داده‌های جدول: ردیف | وسیله | کاربرد / ۱ | پنل خورشیدی | دریافت نور / ۲ | باتری | ذخیره انرژی",
        "جای تصویر: تصویر آموزشی معرفی‌شده توسط مدرس را اینجا درج کنید.",
    ]:
        p = doc.add_paragraph()
        base.set_rtl(p)
        run = p.add_run(value)
        base.style_run(run, size=11)
    target = FILES / "word-grade7-practice-starter.docx"
    doc.save(target)
    return target


def build_sample_project() -> tuple[Path, Path]:
    FILES.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    header = section.header.paragraphs[0]
    base.set_rtl(header)
    run = header.add_run("نمونه پروژه Word پایه هفتم | علوم")
    base.style_run(run, bold=True, size=9, color=BLUE)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.add_page_number(footer)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("چرخه آب در طبیعت")
    base.style_run(run, bold=True, size=22, color=BLUE)
    p = doc.add_paragraph()
    base.set_rtl(p)
    base.add_text_with_markup(p, "آب در طبیعت پیوسته میان زمین و جو حرکت می‌کند. گرمای خورشید آب را تبخیر می‌کند، بخار در ارتفاع سرد و به ابر تبدیل می‌شود و سپس بارش دوباره آب را به زمین بازمی‌گرداند.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(VISUALS / "12-water-cycle.png"), width=Inches(6.2))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("شکل ۱ ـ مراحل اصلی چرخه آب")
    base.style_run(run, size=9, color=GRAY)
    p = doc.add_paragraph()
    base.set_rtl(p)
    run = p.add_run("مراحل چرخه آب")
    base.style_run(run, bold=True, size=15, color=VIOLET)
    for value in ["تبخیر آب با گرمای خورشید", "سردشدن بخار و تشکیل ابر", "بارش باران یا برف", "جمع‌شدن و جریان آب"]:
        p = doc.add_paragraph(style="List Number")
        base.set_rtl(p)
        run = p.add_run(value)
        base.style_run(run, size=11)
    doc.add_page_break()
    p = doc.add_paragraph()
    base.set_rtl(p)
    run = p.add_run("واژه‌های کلیدی")
    base.style_run(run, bold=True, size=16, color=VIOLET)
    table = doc.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    rows = [
        ["فرایند", "تعریف کوتاه", "نمونه"],
        ["تبخیر", "تبدیل آب مایع به بخار", "خشک‌شدن لباس"],
        ["میعان", "تبدیل بخار به قطره", "قطره روی شیشه سرد"],
        ["بارش", "بازگشت آب از ابر", "باران و برف"],
        ["رواناب", "حرکت آب روی زمین", "رودخانه"],
    ]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            p = cell.paragraphs[0]
            base.set_rtl(p)
            run = p.add_run(value)
            base.style_run(run, bold=r == 0, size=10, color="#FFFFFF" if r == 0 else DARK)
            if r == 0:
                base.set_cell_shading(cell, BLUE.removeprefix("#"))
    p = doc.add_paragraph()
    base.set_rtl(p)
    run = p.add_run("فکر کنید")
    base.style_run(run, bold=True, size=15, color=VIOLET)
    p = doc.add_paragraph()
    base.set_rtl(p)
    base.add_text_with_markup(p, "اگر گرمای خورشید کمتر شود، کدام بخش چرخه آب زودتر تغییر می‌کند؟ پاسخ خود را در سه جمله بنویسید.")
    p = doc.add_paragraph()
    base.set_rtl(p)
    base.add_text_with_markup(p, "منبع آموزشی: محتوای نمونه تولیدشده برای تمرین صفحه‌آرایی Word.")
    target = FILES / "word-grade7-sample-project.docx"
    doc.save(target)

    sample_html = OUTPUT / "sample-project.html"
    sample_html.write_text(
        f"""<!doctype html><html lang="fa" dir="rtl"><head><meta charset="utf-8"><style>{html_css()}</style></head>
<body><h1 style="page-break-before:auto">چرخه آب در طبیعت</h1>
<p>آب در طبیعت پیوسته میان زمین و جو حرکت می‌کند. گرمای خورشید آب را تبخیر می‌کند، بخار در ارتفاع سرد و به ابر تبدیل می‌شود و سپس بارش دوباره آب را به زمین بازمی‌گرداند.</p>
<img src="{(VISUALS / '12-water-cycle.png').as_uri()}"><p style="text-align:center">شکل ۱ ـ مراحل اصلی چرخه آب</p>
<h2>مراحل چرخه آب</h2><ol><li>تبخیر آب با گرمای خورشید</li><li>سردشدن بخار و تشکیل ابر</li><li>بارش باران یا برف</li><li>جمع‌شدن و جریان آب</li></ol>
<div style="page-break-before:always"><h1 style="page-break-before:auto">واژه‌های کلیدی</h1>
<table><thead><tr><th>فرایند</th><th>تعریف کوتاه</th><th>نمونه</th></tr></thead><tbody>
<tr><td>تبخیر</td><td>تبدیل آب مایع به بخار</td><td>خشک‌شدن لباس</td></tr>
<tr><td>میعان</td><td>تبدیل بخار به قطره</td><td>قطره روی شیشه سرد</td></tr>
<tr><td>بارش</td><td>بازگشت آب از ابر</td><td>باران و برف</td></tr>
<tr><td>رواناب</td><td>حرکت آب روی زمین</td><td>رودخانه</td></tr></tbody></table>
<h2>فکر کنید</h2><p>اگر گرمای خورشید کمتر شود، کدام بخش چرخه آب زودتر تغییر می‌کند؟ پاسخ خود را در سه جمله بنویسید.</p></div>
</body></html>""",
        encoding="utf-8",
    )
    pdf_target = FILES / "word-grade7-sample-project.pdf"
    print_pdf(sample_html, pdf_target)
    sample_html.unlink()
    return target, pdf_target


def main() -> None:
    if not LOGO.exists():
        raise FileNotFoundError(f"School logo is missing: {LOGO}")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    print("Creating Word lesson visuals...", flush=True)
    create_visuals()
    print("Building student booklet...", flush=True)
    html_path = build_html()
    docx_path = build_docx()
    pdf_path = OUTPUT / "word-grade7-booklet.pdf"
    print_pdf(html_path, pdf_path)
    print("Building companion practice files...", flush=True)
    practice = build_practice_file()
    sample_docx, sample_pdf = build_sample_project()
    for path in (docx_path, pdf_path, html_path, practice, sample_docx, sample_pdf):
        print(f"Generated: {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
